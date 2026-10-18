import csv,io,json,secrets
from datetime import timedelta
from zipfile import ZipFile,ZIP_DEFLATED
import pytest
from openpyxl import Workbook
from sqlalchemy import select,func
from app.models.entities import Product,CatalogItemSource,CatalogImportJob,CatalogSource
from app.services import catalog_jobs,catalog_providers,safe_network
from app.services.catalog_normalize import parse_upload,InvalidCatalog
from app.core.security import utcnow,decrypt_text
from app.core.config import get_settings
from catalog_helpers import attach,verify

def source(owner,kind="file",key="qa-source",**extra):
    response=owner.post("/api/admin/catalog/sources",json={"key":key,"name":"QA authorized source","kind":kind,"owner_authorized":True,**extra})
    assert response.status_code==201,response.text
    return response.json()

def upload(owner,source,rows,request_id=None):
    request_id=request_id or secrets.token_hex(16)
    return owner.post("/api/admin/catalog/uploads",data={"source_id":source["id"],"request_id":request_id},files={"file":("qa.json",json.dumps(rows).encode(),"application/json")})

def upload_file(owner,source,name,raw):
    return owner.post("/api/admin/catalog/uploads",data={"source_id":source["id"],"request_id":secrets.token_hex(16)},files={"file":(name,raw,"application/octet-stream")})

def csv_file(headers,values):
    output=io.StringIO()
    writer=csv.writer(output)
    writer.writerow(headers)
    writer.writerow(values)
    return output.getvalue().encode("utf-8")

def fixture_row(**changes):
    return {"external_id":"qa-001","name":"QA catalog item","brand":"QA only","size":"30 ml","category":"skincare","description":"Isolated QA catalog description.","ingredients":["QA fixture"],"usage":"QA label only.","price":"9999.00","stock":999,"active":True,"content_verified":True,**changes}

def test_uploaded_feed_creates_only_unpriced_drafts(client,make_client,db):
    owner=make_client("super_admin")
    src=source(owner)
    queued=upload(owner,src,[fixture_row()])
    assert queued.status_code==202,queued.text
    assert db.scalar(select(func.count()).select_from(Product))==0
    assert catalog_jobs.process_one(db)==1
    product=db.scalar(select(Product))
    assert product.publication_state=="DRAFT" and not product.active and not product.content_verified
    assert product.price==0 and product.stock==0 and product.acquisition_cost==0
    assert product.image_url is None
    assert client.get("/api/products").json()==[]
    job=db.get(CatalogImportJob,queued.json()["id"])
    assert job.status=="completed" and job.counts_json["created"]==1

def test_queue_request_replay_and_one_active_job_per_source(make_client,db):
    owner=make_client("super_admin")
    src=source(owner)
    ident=secrets.token_hex(16)
    first=upload(owner,src,[fixture_row()],ident)
    assert first.status_code==202
    assert upload(owner,src,[fixture_row()],ident).status_code==202
    assert upload(owner,src,[fixture_row(name="Another payload")],ident).status_code==409
    assert upload(owner,src,[fixture_row()]).status_code==409
    assert db.scalar(select(func.count()).select_from(CatalogImportJob))==1

def test_refresh_keeps_verified_local_price_stock_and_content(make_client,db):
    owner=make_client("super_admin")
    src=source(owner)
    assert upload(owner,src,[fixture_row()]).status_code==202
    catalog_jobs.process_one(db)
    product=db.scalar(select(Product))
    edited=owner.patch("/api/admin/products/"+str(product.id),json={"price":"250.25","stock":5,"acquisition_cost":"120.00","condition":"sealed","expiry_required":False}).json()
    edited,asset=attach(owner,edited)
    published=verify(owner,edited)
    again=upload(owner,src,[fixture_row(name="Changed external title",price=1,stock=10000)])
    assert again.status_code==202,again.text
    catalog_jobs.process_one(db)
    assert all(j.status=="completed" for j in db.scalars(select(CatalogImportJob)).all()), [(j.status,j.error,j.counts_json) for j in db.scalars(select(CatalogImportJob)).all()]
    db.expire_all()
    product=db.get(Product,product.id)
    assert product.price==250.25 and product.stock==5 and product.name=="QA catalog item"
    assert product.active and product.publication_state=="PUBLISHED"
    item=db.scalar(select(CatalogItemSource))
    assert item.changes_json["name"]["after"]=="Changed external title"
    assert product.image_url==asset["url"]

def test_identical_reimport_is_idempotent(make_client,db):
    owner=make_client("super_admin")
    src=source(owner)
    for _ in range(2):
        response=upload(owner,src,[fixture_row()])
        assert response.status_code==202,response.text
        catalog_jobs.process_one(db)
    assert db.scalar(select(func.count()).select_from(Product))==1
    assert db.scalar(select(func.count()).select_from(CatalogItemSource))==1
    latest=db.scalars(select(CatalogImportJob).order_by(CatalogImportJob.created_at.desc())).first()
    assert latest.counts_json["unchanged"]==1, [(j.status,j.error,j.counts_json) for j in db.scalars(select(CatalogImportJob)).all()]

def test_csv_utf8_bom_malicious_fields_and_duplicate_headers(make_client,db):
    owner=make_client("super_admin")
    src=source(owner)
    raw='\ufeffexternal_id,name,brand,price,stock,active,content_verified\ncsv-1,"=2+2",QA,10,999,true,true\n'.encode()
    result=owner.post("/api/admin/catalog/uploads",data={"source_id":src["id"],"request_id":secrets.token_hex(16)},files={"file":("qa.csv",raw,"text/csv")})
    assert result.status_code==202,result.text
    catalog_jobs.process_one(db)
    product=db.scalar(select(Product))
    assert product.name=="=2+2" and product.price==0 and product.stock==0 and not product.content_verified
    bad=owner.post("/api/admin/catalog/uploads",data={"source_id":src["id"],"request_id":secrets.token_hex(16)},files={"file":("bad.csv",b"id,name,name\n1,a,b","text/csv")})
    assert bad.status_code==422

def test_source_authorization_permissions_and_encrypted_header_credentials(make_client,db):
    inventory=make_client("inventory_manager")
    assert inventory.post("/api/admin/catalog/sources",json={"key":"forged-source","name":"Forged","kind":"file","owner_authorized":True}).status_code==403
    owner=make_client("super_admin")
    registered=source(owner,kind="supplier_api",endpoint="https://example.com/products",credential="offline-header-fixture")
    assert "offline-header-fixture" not in str(registered)
    stored=db.get(CatalogSource,registered["id"])
    assert "offline-header-fixture" not in stored.credential_blob
    assert decrypt_text(stored.credential_blob,get_settings().pii_encryption_key)=="offline-header-fixture"

def test_provider_failure_retries_and_resumes_without_duplicate_products(make_client,db,monkeypatch):
    owner=make_client("super_admin")
    src=source(owner,kind="supplier_api",endpoint="https://example.com/products")
    request=owner.post("/api/admin/catalog/jobs",json={"source_id":src["id"],"request_id":secrets.token_hex(16),"pages":1})
    assert request.status_code==202,request.text
    def fail(*args):raise safe_network.ExternalFetchError("HTTP_503",True)
    monkeypatch.setattr(catalog_providers,"fetch_page",fail)
    catalog_jobs.process_one(db)
    job=db.get(CatalogImportJob,request.json()["id"])
    assert job.status=="queued" and job.attempts==1 and job.error=="HTTP_503"
    assert catalog_jobs.process_one(db)==0
    job.next_attempt_at=utcnow()-timedelta(seconds=1)
    db.commit()
    monkeypatch.setattr(catalog_providers,"fetch_page",lambda *args:[fixture_row()])
    assert catalog_jobs.process_one(db)==1
    db.refresh(job)
    assert job.status=="completed" and job.counts_json["created"]==1

def test_partial_scan_does_not_mark_unseen_references_removed(make_client,db,monkeypatch):
    owner=make_client("super_admin")
    src=source(owner,kind="supplier_api",endpoint="https://example.com/products")
    full=[fixture_row(external_id="qa-"+str(i)) for i in range(12)]
    monkeypatch.setattr(catalog_providers,"fetch_page",lambda *args:full)
    request=owner.post("/api/admin/catalog/jobs",json={"source_id":src["id"],"request_id":secrets.token_hex(16),"pages":1})
    assert request.status_code==202
    catalog_jobs.process_one(db)
    monkeypatch.setattr(catalog_providers,"fetch_page",lambda *args:[])
    later=owner.post("/api/admin/catalog/jobs",json={"source_id":src["id"],"request_id":secrets.token_hex(16),"pages":1,"start":12})
    assert later.status_code==202
    catalog_jobs.process_one(db)
    assert all(item.available_at_source for item in db.scalars(select(CatalogItemSource)).all())

def test_feed_cannot_copy_price_stock_or_approval_through_apply(make_client,db):
    owner=make_client("super_admin")
    src=source(owner)
    upload(owner,src,[fixture_row()])
    catalog_jobs.process_one(db)
    item=db.scalar(select(CatalogItemSource))
    product=db.get(Product,item.product_id)
    for field in ["price","stock","active","image_url"]:
        response=owner.post("/api/admin/catalog/items/"+str(item.id)+"/apply",json={"fields":[field],"expected_revision":product.revision,"reason":"Explicit isolated QA change review"})
        assert response.status_code==422

def test_website_csv_preserves_one_unverified_image_candidate_after_worker(make_client,db):
    owner=make_client("super_admin")
    src=source(owner)
    raw=csv_file(["SL","Brand","Product Name","Size","বাংলা Description","Barcode","Product Source Link","Direct Image URL","BD Price ৳","TP ৳","GlowNuture Stock","Image Verification","Ingredients (truncated)"],
        [1,"The Ordinary","Caffeine solution","30 mL","Reference description","0769915195712","https://world.openbeautyfacts.org/product/0769915195712","https://images.openbeautyfacts.org/sample.jpg","999","500","8","Source dataset; not live-checked","Partial ingredients"])
    queued=upload_file(owner,src,"website.csv",raw)
    assert queued.status_code==202,queued.text
    assert catalog_jobs.process_one(db)==1
    product=db.scalar(select(Product))
    item=db.scalar(select(CatalogItemSource))
    assert product.price==0 and product.stock==0 and product.image_url is None
    assert not product.active and product.publication_state=="DRAFT"
    assert item.external_id=="0769915195712"
    assert item.snapshot_json["image_candidates"]==[{"url":"https://images.openbeautyfacts.org/sample.jpg","reported_license":"","rights_verified":False}]
    assert item.snapshot_json["source_metadata"]["image_verification"]=="Source dataset; not live-checked"
    assert item.snapshot_json["reference_prices_bdt"]=={"submitted_bd_price":"999","submitted_trade_price":"500"}
    assert item.snapshot_json["reference_only"] is True

def test_website_csv_500_rows_processes_five_pages_and_reimports_cleanly(make_client,db):
    owner=make_client("super_admin")
    src=source(owner)
    output=io.StringIO()
    writer=csv.writer(output)
    writer.writerow(["Brand","Product Name","Size","Barcode","Direct Image URL"])
    for number in range(500):
        writer.writerow(["Fixture Brand",f"Fixture item {number}","30 mL",str(1000000000000+number),f"https://images.example.com/{number}.jpg"])
    raw=output.getvalue().encode("utf-8")
    for expected_field in ("created","unchanged"):
        queued=upload_file(owner,src,"website.csv",raw)
        assert queued.status_code==202,queued.text
        for _ in range(5):assert catalog_jobs.process_one(db)==1
        job=db.get(CatalogImportJob,queued.json()["id"])
        assert job.status=="completed" and job.counts_json[expected_field]==500
        assert db.scalar(select(func.count()).select_from(Product))==500
        assert db.scalar(select(func.count()).select_from(CatalogItemSource))==500

def test_ordinary_xlsx_uses_first_data_sheet_and_keeps_prices_as_references(make_client,db):
    book=Workbook()
    book.active.title="Empty"
    sheet=book.create_sheet("The Ordinary BD")
    sheet.append(["SL","Product name","Size","Category","বাংলা Description","Retail price (৳)","Listed MRP (৳)","TP / Cost (৳)","GlowNuture Stock","Retailer availability","Retailer source link","Price checked","Notes"])
    sheet.append([1,"Niacinamide 10% + Zinc 1%","30 mL","serum","Reference description",1099,1650,700,4,"Listed - verify stock","https://example.com/product","2026-10-16","Verify label"])
    later=book.create_sheet("Ignore later")
    later.append(["name","sku"])
    later.append(["Wrong product","wrong"])
    output=io.BytesIO()
    book.save(output)
    book.close()
    owner=make_client("super_admin")
    src=source(owner)
    queued=upload_file(owner,src,"ordinary.xlsx",output.getvalue())
    assert queued.status_code==202,queued.text
    assert catalog_jobs.process_one(db)==1
    assert db.scalar(select(func.count()).select_from(Product))==1
    product=db.scalar(select(Product))
    item=db.scalar(select(CatalogItemSource))
    assert product.brand=="The Ordinary" and product.price==0 and product.stock==0
    assert item.external_id.startswith("ordinary-sheet-")
    assert item.snapshot_json["reference_prices_bdt"]=={"listed_retail_price":"1099","listed_mrp":"1650","listed_trade_price":"700"}
    assert item.snapshot_json["url"]=="https://example.com/product"
    assert item.snapshot_json["source_metadata"]["retailer_availability"]=="Listed - verify stock"

def test_shajgoj_csv_refresh_changes_reference_price_only(make_client,db):
    owner=make_client("super_admin")
    src=source(owner)
    headers=["No","brand","product","category","size","mrp_BDT","observed_price_BDT","status","short_description","how_to_use","source_page","source_section","other_page_price_BDT"]
    def feed(price):
        return csv_file(headers,[1,"COSRX","COSRX Acne Pimple Master Patch","Acne Patch","Pack","555.0",price,"Listed - verify stock","Reference description","Use as directed","https://shop.shajgoj.com/product-category/k-beauty","K-Beauty",""])
    assert upload_file(owner,src,"prices.csv",feed("349.0")).status_code==202
    assert catalog_jobs.process_one(db)==1
    product=db.scalar(select(Product))
    saved=owner.patch("/api/admin/products/"+str(product.id),json={"price":"450.00","stock":7})
    assert saved.status_code==200,saved.text
    assert upload_file(owner,src,"prices.csv",feed("399.0")).status_code==202
    assert catalog_jobs.process_one(db)==1
    db.expire_all()
    product=db.get(Product,product.id)
    item=db.scalar(select(CatalogItemSource))
    assert db.scalar(select(func.count()).select_from(Product))==1
    assert product.price==450 and product.stock==7
    assert item.external_id.startswith("shajgoj-")
    assert item.snapshot_json["reference_prices_bdt"]["competitor_observed_price"]=="399.0"
    assert item.changes_json["reference_prices_bdt"]["before"]["competitor_observed_price"]=="349.0"
    assert item.snapshot_json["reference_only"] is True

def test_xlsx_rejects_oversized_uncompressed_archive():
    output=io.BytesIO()
    with ZipFile(output,"w",ZIP_DEFLATED) as archive:
        archive.writestr("xl/worksheets/sheet1.xml",b" "*8_000_001)
    assert len(output.getvalue())<1_000_000
    with pytest.raises(InvalidCatalog,match="safe limits"):
        parse_upload(output.getvalue(),"oversized.xlsx")

def test_bundled_import_requires_owner_and_explicit_authorization(make_client,db):
    payload={"owner_authorized":True,"request_id":secrets.token_hex(16)}
    for role in ("inventory_manager","admin"):
        denied=make_client(role).post("/api/admin/catalog/bundled-imports",json=payload)
        assert denied.status_code==403,denied.text
    owner=make_client("super_admin")
    missing=owner.post("/api/admin/catalog/bundled-imports",json={**payload,"owner_authorized":False})
    assert missing.status_code==422,missing.text
    assert db.scalar(select(func.count()).select_from(CatalogSource))==0
    assert db.scalar(select(func.count()).select_from(CatalogImportJob))==0

def test_bundled_import_creates_739_drafts_and_reimport_is_idempotent(client,make_client,db):
    owner=make_client("super_admin")
    endpoint="/api/admin/catalog/bundled-imports"
    request={"owner_authorized":True,"request_id":secrets.token_hex(16)}
    first=owner.post(endpoint,json=request)
    assert first.status_code==202,first.text
    first_ids={job["id"] for job in first.json()["jobs"]}
    assert len(first_ids)==3
    assert db.scalar(select(func.count()).select_from(CatalogSource))==3
    assert db.scalar(select(func.count()).select_from(CatalogImportJob))==3
    sources={source.key:source for source in db.scalars(select(CatalogSource)).all()}
    assert set(sources)=={"glownuture-website-500","ordinary-bd-reference","shajgoj-price-reference"}
    assert all(source.kind=="file" and source.enabled and source.authorized_by for source in sources.values())

    replay=owner.post(endpoint,json=request)
    assert replay.status_code==202,replay.text
    assert {job["id"] for job in replay.json()["jobs"]}==first_ids
    assert db.scalar(select(func.count()).select_from(CatalogImportJob))==3

    for _ in range(9):assert catalog_jobs.process_one(db)==1
    assert catalog_jobs.process_one(db)==0
    assert db.scalar(select(func.count()).select_from(Product))==739
    assert db.scalar(select(func.count()).select_from(CatalogItemSource))==739
    assert client.get("/api/products").json()==[]
    assert db.scalar(select(func.count()).select_from(Product).where(Product.price!=0))==0
    assert db.scalar(select(func.count()).select_from(Product).where(Product.stock!=0))==0
    assert db.scalar(select(func.count()).select_from(Product).where(Product.active.is_(True)))==0
    assert db.scalar(select(func.count()).select_from(Product).where(Product.image_url.is_not(None)))==0
    assert all(item.snapshot_json["reference_only"] is True for item in db.scalars(select(CatalogItemSource)).all())
    first_jobs=db.scalars(select(CatalogImportJob).where(CatalogImportJob.id.in_(first_ids))).all()
    assert sorted(job.counts_json["created"] for job in first_jobs)==[12,227,500]
    assert all(job.status=="completed" for job in first_jobs)

    later=owner.post(endpoint,json={**request,"request_id":secrets.token_hex(16)})
    assert later.status_code==202,later.text
    later_ids={job["id"] for job in later.json()["jobs"]}
    assert later_ids.isdisjoint(first_ids)
    for _ in range(9):assert catalog_jobs.process_one(db)==1
    assert catalog_jobs.process_one(db)==0
    later_jobs=db.scalars(select(CatalogImportJob).where(CatalogImportJob.id.in_(later_ids))).all()
    assert sorted(job.counts_json["unchanged"] for job in later_jobs)==[12,227,500]
    assert all(job.status=="completed" for job in later_jobs)
    assert db.scalar(select(func.count()).select_from(Product))==739
    assert db.scalar(select(func.count()).select_from(CatalogItemSource))==739
    assert client.get("/api/products").json()==[]
