"""One bounded, plain-text catalog DTO. Source prices/stock never become store inventory."""
import csv
import hashlib
import io
import json
import re
from zipfile import ZipFile
from html import unescape
from html.parser import HTMLParser
from openpyxl import load_workbook
from .safe_network import validate_url,ExternalFetchError

class InvalidCatalog(ValueError):pass
class Plain(HTMLParser):
    def __init__(self):
        super().__init__()
        self.parts=[]
    def handle_data(self,data):self.parts.append(data)

def text(value,maximum=12000):
    if value is None:return ""
    if not isinstance(value,(str,int,float)):return ""
    parser=Plain()
    parser.feed(str(value))
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]","",unescape(" ".join(parser.parts))).strip()[:maximum]

def url(value):
    if not isinstance(value,str):return ""
    try:validate_url(value)
    except (ExternalFetchError,ValueError):return ""
    return value

def array(value):
    if isinstance(value,list):return value.copy()
    if isinstance(value,str):
        if value.strip().startswith("["):
            try:
                parsed=json.loads(value)
                if isinstance(parsed,list):return parsed
            except ValueError:pass
        return [part.strip() for part in re.split(r"[\n;]",value) if part.strip()]
    return []

def gtin(value):
    digits=str(value or "").strip()
    if len(digits) not in {8,12,13,14} or not digits.isascii() or not digits.isdigit():return ""
    total=sum(int(digit)*(3 if index%2==0 else 1) for index,digit in enumerate(reversed(digits[:-1])))
    return digits if (10-total%10)%10==int(digits[-1]) else ""

def _reference_fields(value,allowed,maximum=1000):
    if not isinstance(value,dict):return {}
    return {key:clean for key in allowed if (clean:=text(value.get(key),maximum))}

def _header(value):
    return re.sub(r"\s+"," ",str(value or "").replace("\ufeff","")).strip().casefold()

def _check_headers(headers):
    names=[_header(name) for name in headers]
    if not names or len(names)>60 or not all(names) or len(set(names))!=len(names):
        raise InvalidCatalog("Catalog headers must be unique, named and bounded")

def _stable_id(prefix,brand,name,size):
    parts=[re.sub(r"\s+"," ",text(value,300)).casefold() for value in (brand,name,size)]
    if not parts[1]:raise InvalidCatalog("A product name is required")
    return prefix+"-"+hashlib.sha256("\x1f".join(parts).encode()).hexdigest()[:32]

def _adapt_rows(rows):
    headers={_header(key) for key in rows[0]}
    if {"product name","barcode","direct image url"}<=headers:
        kind="website"
    elif {"product name","retail price (৳)","retailer source link"}<=headers:
        kind="ordinary_sheet"
    elif {"product","observed_price_bdt","source_page"}<=headers:
        kind="shajgoj"
    else:return rows
    adapted=[]
    for row in rows:
        values={_header(key):value for key,value in row.items()}
        get=lambda key:values.get(key,"")
        if kind=="website":
            brand,name,size=get("brand"),get("product name"),get("size")
            barcode=text(get("barcode"),80)
            adapted.append({"external_id":barcode or _stable_id("website",brand,name,size),"barcode":barcode,
                "name":name,"brand":brand,"size":size,"description":get("বাংলা description"),
                "source_url":get("product source link"),"image_url":get("direct image url"),
                "reference_prices_bdt":{"submitted_bd_price":get("bd price ৳"),"submitted_trade_price":get("tp ৳")},
                "source_metadata":{"image_verification":get("image verification"),"ingredients_excerpt":get("ingredients (truncated)")},
                "reference_only":True})
        elif kind=="ordinary_sheet":
            name,size=get("product name"),get("size")
            adapted.append({"external_id":_stable_id("ordinary-sheet","The Ordinary",name,size),
                "name":name,"brand":"The Ordinary","size":size,"category":get("category"),
                "description":get("বাংলা description"),"source_url":get("retailer source link"),
                "reference_prices_bdt":{"listed_retail_price":get("retail price (৳)"),
                    "listed_mrp":get("listed mrp (৳)"),"listed_trade_price":get("tp / cost (৳)")},
                "source_metadata":{"retailer_availability":get("retailer availability"),
                    "price_checked":get("price checked"),"notes":get("notes")},
                "reference_only":True})
        else:
            brand,name,size=get("brand"),get("product"),get("size")
            adapted.append({"external_id":_stable_id("shajgoj",brand,name,size),
                "name":name,"brand":brand,"size":size,"category":get("category"),
                "short_description":get("short_description"),"usage":get("how_to_use"),
                "source_url":get("source_page"),
                "reference_prices_bdt":{"competitor_observed_price":get("observed_price_bdt"),
                    "competitor_mrp":get("mrp_bdt"),"competitor_other_page_price":get("other_page_price_bdt")},
                "source_metadata":{"retailer_status":get("status"),"source_section":get("source_section")},
                "reference_only":True})
    return adapted

def _xlsx_rows(raw):
    try:
        with ZipFile(io.BytesIO(raw)) as archive:
            entries=archive.infolist()
            if len(entries)>200 or sum(entry.file_size for entry in entries)>8_000_000 or any(entry.flag_bits&1 for entry in entries):
                raise InvalidCatalog("XLSX archive exceeds safe limits")
        book=load_workbook(io.BytesIO(raw),read_only=True,data_only=True,keep_links=False)
        try:
            for sheet in book.worksheets:
                if sheet.max_row is None or sheet.max_column is None:
                    sheet.calculate_dimension(force=True)
                if sheet.max_column>60 or sheet.max_row>502:
                    raise InvalidCatalog("XLSX maximum is 60 columns and 500 products")
                iterator=sheet.iter_rows(max_row=sheet.max_row,max_col=sheet.max_column,values_only=True)
                first=next((row for row in iterator if sum(value is not None and str(value).strip()!="" for value in row)>1),None)
                if first is None:continue
                headers=list(first)
                while headers and (headers[-1] is None or str(headers[-1]).strip()==""):headers.pop()
                _check_headers(headers)
                rows=[]
                for values in iterator:
                    if not any(value is not None and str(value).strip()!="" for value in values):continue
                    if any(value is not None and str(value).strip()!="" for value in values[len(headers):]):
                        raise InvalidCatalog("XLSX contains values outside its headers")
                    rows.append(dict(zip(headers,values)))
                    if len(rows)>500:raise InvalidCatalog("Maximum 500 products per file")
                return rows
            raise InvalidCatalog("XLSX has no product sheet")
        finally:book.close()
    except InvalidCatalog:raise
    except Exception:raise InvalidCatalog("Malformed XLSX catalog file") from None

def normalize(row,source):
    if not isinstance(row,dict):raise InvalidCatalog("Product row must be an object")
    barcode=gtin(row.get("gtin") or row.get("barcode") or row.get("code"))
    external=row.get("external_id") or row.get("id") or row.get("product_id") or row.get("sku") or barcode or row.get("url")
    name=text(row.get("name_bn") or row.get("name") or row.get("title") or row.get("product_name") or row.get("name_en"),300)
    if not external and source.kind=="ordinary":external=name
    if not isinstance(external,(str,int)) or not str(external).strip() or len(str(external))>200 or not name:
        raise InvalidCatalog("Stable external ID/SKU/GTIN and a product name are required")
    images=[]
    raw=array(row.get("image_candidates") or row.get("images"))
    raw.extend([row.get("image_url"),row.get("image_front_url")])
    for image in raw:
        image=image if isinstance(image,dict) else {"url":image}
        candidate=url(image.get("url") or image.get("src"))
        if candidate and candidate not in [item["url"] for item in images]:
            images.append({"url":candidate,"reported_license":text(image.get("reported_license") or image.get("license"),200),"rights_verified":False})
    ingredients=array(row.get("key_ingredients") or row.get("ingredients") or row.get("ingredients_text") or row.get("ingredients_text_en"))
    ingredients=[text(item.get("name") if isinstance(item,dict) else item,500) for item in ingredients][:80]
    variants=[]
    for item in array(row.get("variants"))[:30]:
        if isinstance(item,dict):
            variants.append({"external_id":text(item.get("external_id") or item.get("id") or item.get("sku"),200),"sku":text(item.get("sku"),100),"label":text(item.get("label") or item.get("size") or item.get("name"),160),"size":text(item.get("size"),80),"gtin":gtin(item.get("gtin") or item.get("barcode"))})
    return {"provider":source.key,"external_id":str(external).strip(),"gtin":barcode,"name":name,"en":text(row.get("name_en") or row.get("en") or row.get("product_name_en"),300),"brand":text(row.get("brand") or row.get("brands"),120),"size":text(row.get("size") or row.get("quantity"),80),"unit":text(row.get("unit"),30),"category":text(row.get("category") or row.get("categories"),80),"description":text(row.get("description") or row.get("long_description")),"short":text(row.get("short_description") or row.get("short"),1000),"ingredients":list(filter(None,ingredients)),"usage":text(row.get("instructions") or row.get("usage") or row.get("directions"),3000),"warnings":text(row.get("warnings"),3000),"origin":text(row.get("origin") or row.get("country_of_origin"),120),"url":url(row.get("source_url") or row.get("url")),"image_candidates":images[:8],"variants":variants,"attribution":source.attribution,"license":source.license,"reference_only":source.kind=="openbeautyfacts" or row.get("reference_only") is True,"reference_prices_bdt":_reference_fields(row.get("reference_prices_bdt"),{"submitted_bd_price","submitted_trade_price","listed_retail_price","listed_mrp","listed_trade_price","competitor_observed_price","competitor_mrp","competitor_other_page_price"},60),"source_metadata":_reference_fields(row.get("source_metadata"),{"image_verification","ingredients_excerpt","retailer_availability","price_checked","notes","retailer_status","source_section"})}

def rows_from_json(payload):
    data=payload.get("data",payload) if isinstance(payload,dict) else payload
    if isinstance(data,dict):data=data.get("products",data.get("items"))
    if not isinstance(data,list):raise InvalidCatalog("Expected a product array in data.products, data, products or items")
    return data

def parse_upload(raw,filename):
    if len(raw)>1_000_000:raise InvalidCatalog("Feed maximum size is 1 MB")
    if filename.lower().endswith(".xlsx"):
        rows=_xlsx_rows(raw)
    else:
        try:value=raw.decode("utf-8-sig")
        except UnicodeDecodeError:raise InvalidCatalog("Use a UTF-8 CSV or JSON file") from None
        try:
            if filename.lower().endswith(".json"):rows=rows_from_json(json.loads(value))
            elif filename.lower().endswith(".csv"):
                csv.field_size_limit(20000)
                reader=csv.DictReader(io.StringIO(value))
                _check_headers(reader.fieldnames or [])
                rows=[]
                for row in reader:
                    if None in row:raise InvalidCatalog("CSV contains extra columns")
                    rows.append(row)
                    if len(rows)>500:raise InvalidCatalog("Maximum 500 products per file")
            else:raise InvalidCatalog("Use a .csv, .json or .xlsx catalog file")
        except (ValueError,csv.Error,RecursionError) as error:
            if isinstance(error,InvalidCatalog):raise
            raise InvalidCatalog("Malformed catalog file") from None
    if not rows or len(rows)>500:raise InvalidCatalog("Supply between 1 and 500 products")
    return _adapt_rows(rows) if filename.lower().endswith((".csv",".xlsx")) else rows

def digest(snapshot):
    return hashlib.sha256(json.dumps(snapshot,sort_keys=True,ensure_ascii=False,separators=(",",":")).encode()).hexdigest()
